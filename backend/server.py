from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

import io
import csv
import uuid
import ssl
import ftplib
import asyncio
import secrets
import logging
import subprocess
import urllib.parse
from datetime import datetime, timezone, timedelta
from typing import List, Optional, Annotated
from email.message import EmailMessage

import bcrypt
import jwt
import requests
import aiosmtplib
from bson import ObjectId
from fastapi import FastAPI, APIRouter, HTTPException, Depends, UploadFile, File, Request, Response, Query
from fastapi.responses import RedirectResponse, FileResponse
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, BeforeValidator, ConfigDict, EmailStr
from PIL import Image, ImageOps

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ['JWT_SECRET']
JWT_ALGORITHM = "HS256"

STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_NAME = "sigflow"

BACKEND_PUBLIC_URL = os.environ.get("REACT_APP_BACKEND_URL", "")
_RESTART_COMMAND = os.environ.get("RESTART_COMMAND", "sudo supervisorctl restart frontend backend")

STORAGE_BACKEND = (os.environ.get("STORAGE_BACKEND") or "emergent").strip().lower()
LOCAL_STORAGE_DIR = Path(os.environ.get("LOCAL_STORAGE_DIR") or (ROOT_DIR / "storage"))

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

app = FastAPI()
api_router = APIRouter(prefix="/api")
security = HTTPBearer(auto_error=False)

# ---------------------------------------------------------------------------
# Object storage helpers — pluggable backend (local disk or Emergent)
# ---------------------------------------------------------------------------
storage_key = None


def init_storage(force: bool = False):
    global storage_key
    if storage_key and not force:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def _emergent_put(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(f"{STORAGE_URL}/objects/{path}",
                        headers={"X-Storage-Key": key, "Content-Type": content_type},
                        data=data, timeout=120)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.put(f"{STORAGE_URL}/objects/{path}",
                            headers={"X-Storage-Key": key, "Content-Type": content_type},
                            data=data, timeout=120)
    resp.raise_for_status()
    return resp.json()


def _emergent_get(path: str):
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")


def _local_path(path: str) -> Path:
    safe = path.replace("..", "").lstrip("/")
    return LOCAL_STORAGE_DIR / safe


def put_object(path: str, data: bytes, content_type: str) -> dict:
    if STORAGE_BACKEND == "local":
        fp = _local_path(path)
        fp.parent.mkdir(parents=True, exist_ok=True)
        fp.write_bytes(data)
        res = {"path": path}
    else:
        res = _emergent_put(path, data, content_type)
    # Publication automatique vers l'hébergement externe (SiteGround FTP ou Cloudflare R2) si activé.
    if _assets_cache.get("enabled"):
        try:
            _publish_asset(_assets_cache, path, data, content_type)
        except Exception as e:
            logger.error(f"Publication hébergement externe échouée pour {path}: {e}")
    return res


# ---------------------------------------------------------------------------
# Hébergement externe des images (FTP/FTPS) — pour des URLs toujours actives,
# indépendantes du backend (ex. sous-domaine SiteGround en HTTPS).
# ---------------------------------------------------------------------------
_assets_cache: dict = {}


async def refresh_assets_cache():
    global _assets_cache
    _assets_cache = await db.config.find_one({"key": "assets"}) or {}


def _ftp_open(cfg: dict):
    host = cfg["ftp_host"]; user = cfg["ftp_user"]; pwd = cfg.get("ftp_password", "")
    if cfg.get("ftp_secure", True):
        ftp = ftplib.FTP_TLS(context=ssl.create_default_context())
        ftp.connect(host, int(cfg.get("ftp_port") or 21), timeout=30)
        ftp.login(user, pwd)
        ftp.prot_p()
    else:
        ftp = ftplib.FTP()
        ftp.connect(host, int(cfg.get("ftp_port") or 21), timeout=30)
        ftp.login(user, pwd)
    return ftp


def _ftp_makedirs(ftp, remote_dir: str):
    parts = [p for p in remote_dir.split("/") if p]
    cur = ""
    for p in parts:
        cur += "/" + p
        try:
            ftp.mkd(cur)
        except Exception:
            pass


def _ftp_upload(cfg: dict, path: str, data: bytes):
    base = (cfg.get("ftp_dir") or "").rstrip("/")
    remote = f"{base}/{path}".lstrip("/") if base else path
    remote = "/" + remote if not remote.startswith("/") else remote
    ftp = _ftp_open(cfg)
    try:
        _ftp_makedirs(ftp, remote.rsplit("/", 1)[0])
        ftp.storbinary(f"STOR {remote}", io.BytesIO(data))
    finally:
        try:
            ftp.quit()
        except Exception:
            pass


def _ftp_delete(cfg: dict, path: str):
    base = (cfg.get("ftp_dir") or "").rstrip("/")
    remote = f"{base}/{path}".lstrip("/") if base else path
    remote = "/" + remote if not remote.startswith("/") else remote
    ftp = _ftp_open(cfg)
    try:
        ftp.delete(remote)
    finally:
        try:
            ftp.quit()
        except Exception:
            pass


def _r2_client(cfg: dict):
    import boto3
    from botocore.config import Config as _BotoConfig
    endpoint = (cfg.get("r2_endpoint") or "").strip() or f"https://{cfg.get('r2_account_id', '').strip()}.r2.cloudflarestorage.com"
    return boto3.client(
        "s3",
        endpoint_url=endpoint,
        aws_access_key_id=cfg.get("r2_access_key"),
        aws_secret_access_key=cfg.get("r2_secret_key"),
        region_name="auto",
        config=_BotoConfig(signature_version="s3v4"),
    )


def _r2_upload(cfg: dict, path: str, data: bytes, content_type: str = "application/octet-stream"):
    _r2_client(cfg).put_object(
        Bucket=cfg["r2_bucket"], Key=path, Body=data,
        ContentType=content_type or "application/octet-stream",
    )


def _r2_list(cfg: dict):
    client = _r2_client(cfg)
    objs, token = [], None
    while True:
        kw = {"Bucket": cfg["r2_bucket"], "MaxKeys": 1000}
        if token:
            kw["ContinuationToken"] = token
        resp = client.list_objects_v2(**kw)
        for o in resp.get("Contents", []):
            objs.append({"key": o["Key"], "size": int(o.get("Size", 0))})
        if resp.get("IsTruncated"):
            token = resp.get("NextContinuationToken")
        else:
            break
    return objs


def _r2_delete(cfg: dict, keys: list):
    if not keys:
        return 0
    client = _r2_client(cfg)
    for i in range(0, len(keys), 1000):
        batch = keys[i:i + 1000]
        client.delete_objects(Bucket=cfg["r2_bucket"], Delete={"Objects": [{"Key": k} for k in batch]})
    return len(keys)


def _guess_ct(path: str) -> str:
    ext = path.rsplit(".", 1)[-1].lower() if "." in path else ""
    return {
        "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif",
        "webp": "image/webp", "svg": "image/svg+xml", "txt": "text/plain",
    }.get(ext, "application/octet-stream")


def _publish_asset(cfg: dict, path: str, data: bytes, content_type: str = None):
    """Route la publication vers Cloudflare R2 ou FTP selon le fournisseur configuré.
    Si R2 est principal et le backup FTP est activé, copie aussi vers SiteGround (best-effort)."""
    ct = content_type or _guess_ct(path)
    if (cfg.get("provider") or "ftp") == "r2":
        _r2_upload(cfg, path, data, ct)
        if cfg.get("backup_enabled") and cfg.get("ftp_host") and cfg.get("ftp_user"):
            try:
                _ftp_upload(cfg, path, data)
            except Exception as e:
                logger.warning(f"Backup FTP (SiteGround) échoué pour {path}: {e}")
    else:
        _ftp_upload(cfg, path, data)


def _rewrite_asset_urls(html: str) -> str:
    """Remplace les URLs d'images du backend par l'URL publique d'hébergement (toujours active)."""
    cfg = _assets_cache
    if not cfg.get("enabled") or not cfg.get("public_url"):
        return html
    pub = cfg["public_url"].rstrip("/")
    bk = (BACKEND_PUBLIC_URL or "").rstrip("/")
    if bk:
        html = html.replace(f"{bk}/api/social-icons/", f"{pub}/social/")
        html = html.replace(f"{bk}/api/files/", f"{pub}/")
    return html


def get_object(path: str):
    if STORAGE_BACKEND == "local":
        fp = _local_path(path)
        if fp.exists():
            ext = fp.suffix.lstrip(".").lower()
            return fp.read_bytes(), MIME_TYPES.get(ext, "application/octet-stream")
        # fallback to Emergent for objects uploaded before switching to local
        if EMERGENT_KEY:
            try:
                return _emergent_get(path)
            except Exception:
                pass
        raise FileNotFoundError(path)
    return _emergent_get(path)


MIME_TYPES = {
    "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
    "gif": "image/gif", "webp": "image/webp", "svg": "image/svg+xml",
}

# ---------------------------------------------------------------------------
# Auth helpers
# ---------------------------------------------------------------------------
def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))


def create_access_token(user_id: str, email: str, role: str) -> str:
    payload = {
        "sub": user_id, "email": email, "role": role,
        "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "access",
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


async def get_current_user(request: Request, creds: Optional[HTTPAuthorizationCredentials] = Depends(security)) -> dict:
    token = creds.credentials if creds else request.cookies.get("access_token")
    if not token:
        raise HTTPException(status_code=401, detail="Non authentifié")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(status_code=401, detail="Utilisateur introuvable")
        user["id"] = str(user["_id"])
        user.pop("_id", None)
        user.pop("password_hash", None)
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Session expirée")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Jeton invalide")


async def require_admin(user: dict = Depends(get_current_user)) -> dict:
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Accès réservé aux administrateurs")
    return user


# Sections d'administration pouvant être accordées individuellement à un utilisateur.
SECTION_PERMS = ["setup", "brand", "gif", "employees", "email", "m365", "deploy", "hosting", "analytics", "install", "system"]


def require_perm(*perms):
    """Autorise l'accès si l'utilisateur est admin OU possède au moins une des permissions demandées."""
    async def _dep(user: dict = Depends(get_current_user)) -> dict:
        if user.get("role") == "admin":
            return user
        if set(user.get("permissions") or []).intersection(perms):
            return user
        raise HTTPException(status_code=403, detail="Accès non autorisé pour cette section")
    return _dep


async def require_any_management(user: dict = Depends(get_current_user)) -> dict:
    """Autorise si admin ou possède au moins une permission de section (lecture transverse, ex. liste des employés)."""
    if user.get("role") == "admin" or (user.get("permissions") or []):
        return user
    raise HTTPException(status_code=403, detail="Accès non autorisé")


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------
PyObjectId = Annotated[str, BeforeValidator(str)]


class SocialLinks(BaseModel):
    linkedin: str = ""
    twitter: str = ""
    facebook: str = ""
    instagram: str = ""
    youtube: str = ""
    tiktok: str = ""


class DeptBanner(BaseModel):
    id: str
    name: str
    gif_images: List[str] = Field(default_factory=list)
    gif_interval_ms: int = 2500
    gif_url: str = ""
    gif_version: int = 0
    banner_link: str = ""
    layout: str = ""


class CompanySettings(BaseModel):
    model_config = ConfigDict(extra="ignore")
    company_name: str = "Mon Entreprise"
    logo_url: str = ""
    website: str = ""
    phone_main: str = ""
    address: str = ""
    disclaimer: str = ""
    primary_color: str = "#2563EB"
    signature_layout: str = "classic"
    logo_width: int = 86
    banner_width: int = 600
    social: SocialLinks = Field(default_factory=SocialLinks)
    social_style: str = "icons"
    gif_images: List[str] = Field(default_factory=list)
    gif_interval_ms: int = 2500
    gif_url: str = ""
    gif_version: int = 0
    banner_link: str = ""
    name_top_border: bool = False
    name_top_border_color: str = "#111827"
    name_top_border_width: int = 2
    logo_position: str = "left"
    department_banners: List[DeptBanner] = Field(default_factory=list)
    typography: dict = Field(default_factory=dict)


class SmtpConfigInput(BaseModel):
    host: str = ""
    port: int = 587
    username: str = ""
    password: Optional[str] = None
    from_email: str = ""
    from_name: str = ""


class DeptCreate(BaseModel):
    name: str


class EmailTestInput(BaseModel):
    to: EmailStr


class UserPublic(BaseModel):
    id: str
    email: str
    name: str
    role: str
    title: str = ""
    phone_ext: str = ""
    direct_line: str = ""
    department: str = ""
    avatar_url: str = ""
    booking_url: str = ""


class RegisterInput(BaseModel):
    email: EmailStr
    password: str
    name: str


class LoginInput(BaseModel):
    email: EmailStr
    password: str


class EmployeeCreate(BaseModel):
    email: EmailStr
    password: str
    name: str
    title: str = ""
    phone_ext: str = ""
    direct_line: str = ""
    department: str = ""


class ProfileUpdate(BaseModel):
    name: Optional[str] = None
    title: Optional[str] = None
    phone_ext: Optional[str] = None
    direct_line: Optional[str] = None
    department: Optional[str] = None
    avatar_url: Optional[str] = None
    avatar_width: Optional[int] = None
    booking_url: Optional[str] = None


class GifGenerateInput(BaseModel):
    images: List[str]
    interval_ms: int = 2500
    department_id: Optional[str] = None
    banner_link: Optional[str] = None
    layout: Optional[str] = None


def _oid(v: str) -> ObjectId:
    try:
        return ObjectId(v)
    except Exception:
        raise HTTPException(status_code=400, detail="Identifiant invalide")


def user_to_public(u: dict) -> dict:
    return {
        "id": str(u.get("id") or u.get("_id")),
        "email": u.get("email", ""),
        "name": u.get("name", ""),
        "role": u.get("role", "employee"),
        "permissions": u.get("permissions") or [],
        "title": u.get("title", ""),
        "phone_ext": u.get("phone_ext", ""),
        "direct_line": u.get("direct_line", ""),
        "department": u.get("department", ""),
        "avatar_url": u.get("avatar_url", ""),
        "avatar_width": int(u.get("avatar_width") or 0),
        "booking_url": u.get("booking_url", ""),
        "m365_linked": bool(u.get("m365_linked", False)),
        "m365_signature_html": u.get("m365_signature_html", ""),
        "m365_synced_at": u.get("m365_synced_at"),
        "signature_installed": bool(u.get("signature_installed", False)),
        "installed_at": u.get("installed_at"),
        "last_sent_at": u.get("last_sent_at"),
        "last_send_status": u.get("last_send_status"),
    }


# ---------------------------------------------------------------------------
# Auth routes
# ---------------------------------------------------------------------------
@api_router.post("/auth/register")
async def register(data: RegisterInput):
    email = data.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Cet email existe déjà")
    doc = {
        "email": email, "password_hash": hash_password(data.password), "name": data.name,
        "role": "employee", "title": "", "phone_ext": "", "direct_line": "",
        "department": "", "avatar_url": "", "created_at": datetime.now(timezone.utc).isoformat(),
    }
    res = await db.users.insert_one(doc)
    doc["id"] = str(res.inserted_id)
    token = create_access_token(doc["id"], email, "employee")
    return {"token": token, "user": user_to_public(doc)}


@api_router.post("/auth/login")
async def login(data: LoginInput):
    email = data.email.lower()
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(data.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Email ou mot de passe invalide")
    user["id"] = str(user["_id"])
    token = create_access_token(user["id"], email, user["role"])
    return {"token": token, "user": user_to_public(user)}


@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user_to_public(user)


@api_router.put("/auth/me")
async def update_me(data: ProfileUpdate, user: dict = Depends(get_current_user)):
    update = {k: v for k, v in data.model_dump().items() if v is not None}
    if update:
        await db.users.update_one({"_id": ObjectId(user["id"])}, {"$set": update})
    fresh = await db.users.find_one({"_id": ObjectId(user["id"])})
    fresh["id"] = str(fresh["_id"])
    return user_to_public(fresh)


# ---------------------------------------------------------------------------
# Settings routes
# ---------------------------------------------------------------------------
async def load_settings() -> dict:
    doc = await db.settings.find_one({"key": "global"})
    if not doc:
        default = CompanySettings().model_dump()
        default["key"] = "global"
        await db.settings.insert_one(default)
        return CompanySettings().model_dump()
    return CompanySettings(**doc).model_dump()


@api_router.get("/settings")
async def get_settings(user: dict = Depends(get_current_user)):
    return await load_settings()


@api_router.put("/settings")
async def update_settings(data: CompanySettings, admin: dict = Depends(require_perm("brand", "setup"))):
    payload = data.model_dump()
    payload["disclaimer"] = _clean_html(payload.get("disclaimer") or "")
    await db.settings.update_one({"key": "global"}, {"$set": payload}, upsert=True)
    await _maybe_auto_cleanup()
    return await load_settings()


# ---------------------------------------------------------------------------
# Upload + file serving
# ---------------------------------------------------------------------------
@api_router.post("/upload")
async def upload(file: UploadFile = File(...), admin: dict = Depends(require_perm("brand", "gif", "setup"))):
    ext = (file.filename.rsplit(".", 1)[-1] if "." in file.filename else "bin").lower()
    path = f"{APP_NAME}/assets/{uuid.uuid4()}.{ext}"
    data = await file.read()
    ctype = file.content_type or MIME_TYPES.get(ext, "application/octet-stream")
    result = put_object(path, data, ctype)
    stored = result["path"]
    return {"path": stored, "url": f"{BACKEND_PUBLIC_URL}/api/files/{stored}"}


@api_router.post("/upload-avatar")
async def upload_avatar(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    ext = (file.filename.rsplit(".", 1)[-1] if "." in file.filename else "bin").lower()
    path = f"{APP_NAME}/avatars/{user['id']}/{uuid.uuid4()}.{ext}"
    data = await file.read()
    ctype = file.content_type or MIME_TYPES.get(ext, "application/octet-stream")
    result = put_object(path, data, ctype)
    stored = result["path"]
    url = f"{BACKEND_PUBLIC_URL}/api/files/{stored}"
    await db.users.update_one({"_id": ObjectId(user["id"])}, {"$set": {"avatar_url": url}})
    await _maybe_auto_cleanup()
    return {"path": stored, "url": url}


ASSETS_DIR = Path(__file__).parent / "assets"


@app.get("/api/social-icons/{name}")
async def social_icon(name: str):
    safe = name.replace("/", "").replace("..", "")
    fp = ASSETS_DIR / "social" / safe
    if not fp.suffix:
        fp = fp.with_suffix(".png")
    if not fp.exists():
        raise HTTPException(status_code=404, detail="Icône introuvable")
    return FileResponse(str(fp), media_type="image/png",
                        headers={"Cache-Control": "public, max-age=86400"})


@app.get("/api/files/{path:path}")
async def serve_file(path: str):
    try:
        data, ctype = get_object(path)
    except Exception:
        raise HTTPException(status_code=404, detail="Fichier introuvable")
    ext = path.rsplit(".", 1)[-1].lower() if "." in path else ""
    return Response(content=data, media_type=MIME_TYPES.get(ext, ctype),
                    headers={"Cache-Control": "public, max-age=300"})


# ---------------------------------------------------------------------------
# Animated GIF generation
# ---------------------------------------------------------------------------
TARGET_W, TARGET_H = 600, 170


@api_router.post("/gif/generate")
async def generate_gif(data: GifGenerateInput, admin: dict = Depends(require_perm("gif", "setup"))):
    if len(data.images) < 1:
        raise HTTPException(status_code=400, detail="Ajoutez au moins une image")
    frames = []
    for src in data.images:
        try:
            content, _ = get_object(src)
            im = Image.open(io.BytesIO(content)).convert("RGB")
            im = ImageOps.fit(im, (TARGET_W, TARGET_H), Image.LANCZOS)
            frames.append(im)
        except Exception as e:
            logger.error(f"Frame load failed for {src}: {e}")
            raise HTTPException(status_code=400, detail=f"Image illisible: {src}")

    buf = io.BytesIO()
    duration = max(300, int(data.interval_ms))
    frames[0].save(
        buf, format="GIF", save_all=True,
        append_images=frames[1:], duration=duration, loop=0, disposal=2, optimize=True,
    )
    gif_bytes = buf.getvalue()

    settings = await load_settings()

    if data.department_id:
        banners = settings.get("department_banners", [])
        target = next((b for b in banners if b["id"] == data.department_id), None)
        if not target:
            raise HTTPException(status_code=404, detail="Département introuvable")
        gif_path = f"{APP_NAME}/brand/banner_{data.department_id}.gif"
        put_object(gif_path, gif_bytes, "image/gif")
        version = target.get("gif_version", 0) + 1
        gif_url = f"{BACKEND_PUBLIC_URL}/api/files/{gif_path}?v={version}"
        target.update({"gif_images": data.images, "gif_interval_ms": duration,
                       "gif_url": gif_url, "gif_version": version})
        if data.banner_link is not None:
            target["banner_link"] = data.banner_link
        if data.layout is not None:
            target["layout"] = data.layout
        await db.settings.update_one({"key": "global"}, {"$set": {"department_banners": banners}})
        return {"gif_url": gif_url, "version": version, "frames": len(frames)}

    gif_path = f"{APP_NAME}/brand/banner.gif"
    put_object(gif_path, gif_bytes, "image/gif")
    version = settings.get("gif_version", 0) + 1
    gif_url = f"{BACKEND_PUBLIC_URL}/api/files/{gif_path}?v={version}"
    update = {"gif_images": data.images, "gif_interval_ms": duration,
              "gif_url": gif_url, "gif_version": version}
    if data.banner_link is not None:
        update["banner_link"] = data.banner_link
    if data.layout is not None:
        update["signature_layout"] = data.layout
    await db.settings.update_one({"key": "global"}, {"$set": update}, upsert=True)
    return {"gif_url": gif_url, "version": version, "frames": len(frames)}


# ---------------------------------------------------------------------------
# Server-side signature builder (mirror of frontend, used for emails)
# ---------------------------------------------------------------------------
def _esc(v):
    return (str(v or "").replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;").replace('"', "&quot;"))


_ALLOWED_HTML_TAGS = ["b", "strong", "i", "em", "u", "br", "p", "div", "span"]


def _clean_html(v):
    """Sanitize admin-authored rich text (disclaimer) to a safe inline tag whitelist."""
    import bleach
    return bleach.clean(str(v or ""), tags=_ALLOWED_HTML_TAGS, attributes={}, strip=True)


def _norm_url(u):
    t = str(u or "").strip()
    if not t:
        return ""
    return t if t.lower().startswith(("http://", "https://")) else f"https://{t}"


def _hex_to_rgb(h):
    x = str(h or "").replace("#", "").strip()
    if len(x) == 3:
        x = "".join(c + c for c in x)
    if len(x) != 6:
        return (37, 99, 235)
    return (int(x[0:2], 16), int(x[2:4], 16), int(x[4:6], 16))


def _dark_accent_for(primary):
    r, g, b = _hex_to_rgb(primary)
    lum = (0.299 * r + 0.587 * g + 0.114 * b) / 255
    if lum < 0.55:
        f = 0.6
        return f"rgb({round(r + (255 - r) * f)},{round(g + (255 - g) * f)},{round(b + (255 - b) * f)})"
    return primary


def _dark_style_block(accent):
    r = (".sf-name{color:#f8fafc!important}.sf-text{color:#e5e7eb!important}.sf-text a{color:#e5e7eb!important}"
         ".sf-muted{color:#9ca3af!important}.sf-disc{color:#94a3b8!important}.sf-disc-line{border-top-color:#334155!important}"
         f".sf-accent{{color:{accent}!important}}.sf-social{{color:#e5e7eb!important}}"
         f".sf-bar{{border-right-color:{accent}!important}}.sf-bar-left{{border-left-color:{accent}!important}}")
    scoped = r.replace(".sf-", ".sf-dark .sf-")
    return f"<style>@media (prefers-color-scheme:dark){{{r}}}{scoped}</style>"


_SOCIAL = {"linkedin": ("LinkedIn", "#0A66C2"), "twitter": ("X", "#111827"),
           "facebook": ("Facebook", "#1877F2"), "instagram": ("Instagram", "#E1306C"),
           "youtube": ("YouTube", "#FF0000"), "tiktok": ("TikTok", "#111827")}


def _resolve_banner(user, s):
    dept = (user.get("department") or "").strip().lower()
    default_layout = s.get("signature_layout") or "classic"
    for b in s.get("department_banners", []):
        if (b.get("name") or "").strip().lower() == dept:
            return {"gif_url": b.get("gif_url", "") or s.get("gif_url", ""),
                    "banner_link": b.get("banner_link", "") or s.get("banner_link", ""),
                    "key": b.get("name"), "layout": b.get("layout") or default_layout}
    return {"gif_url": s.get("gif_url", ""), "banner_link": s.get("banner_link", ""),
            "key": "Défaut", "layout": default_layout}


_FONT_STACKS = {
    "Arial": "Arial, Helvetica, sans-serif",
    "Helvetica": "Helvetica, Arial, sans-serif",
    "Georgia": "Georgia, 'Times New Roman', serif",
    "Times New Roman": "'Times New Roman', Times, serif",
    "Verdana": "Verdana, Geneva, sans-serif",
    "Tahoma": "Tahoma, Geneva, sans-serif",
    "Trebuchet MS": "'Trebuchet MS', Helvetica, sans-serif",
    "Courier New": "'Courier New', Courier, monospace",
    "Calibri": "Calibri, Candara, Segoe, 'Segoe UI', Arial, sans-serif",
}
_TYPO_DEFAULTS = {
    "name": {"font": "Arial", "size": 18, "bold": True, "italic": False, "underline": False, "color": "#0f172a"},
    "title": {"font": "Arial", "size": 13, "bold": True, "italic": False, "underline": False, "color": ""},
    "contact": {"font": "Arial", "size": 13, "bold": False, "italic": False, "underline": False, "color": "#1f2937"},
    "tel": {"font": "Arial", "size": 13, "bold": False, "italic": False, "underline": False, "color": "#1f2937"},
    "direct": {"font": "Arial", "size": 13, "bold": False, "italic": False, "underline": False, "color": "#1f2937"},
    "courriel": {"font": "Arial", "size": 13, "bold": False, "italic": False, "underline": False, "color": "#1f2937"},
    "website": {"font": "Arial", "size": 13, "bold": False, "italic": False, "underline": False, "color": "#1f2937"},
    "address": {"font": "Arial", "size": 12, "bold": False, "italic": False, "underline": False, "color": "#6b7280"},
}


def _typo_style_contact(s, sub, default_color):
    """Style d'une ligne de coordonnées : hérite de 'contact' puis applique le sous-groupe (tel/direct/courriel/website)."""
    d = dict(_TYPO_DEFAULTS["contact"])
    for grp in ("contact", sub):
        cfg = ((s.get("typography") or {}).get(grp)) or {}
        for k, v in cfg.items():
            if v is not None and v != "":
                d[k] = v
    fam = _FONT_STACKS.get(d.get("font") or "Arial", _FONT_STACKS["Arial"])
    col = d.get("color") or default_color
    size = int(d.get("size") or 13)
    css = f"font-family:{fam};font-size:{size}px;color:{col};font-weight:{'700' if d.get('bold') else '400'};"
    if d.get("italic"):
        css += "font-style:italic;"
    if d.get("underline"):
        css += "text-decoration:underline;"
    sub_color = (((s.get("typography") or {}).get(sub)) or {}).get("color")
    return {"css": css, "color": col, "label_color": sub_color or None}


def _typo_color(s, group, default):
    cfg = ((s.get("typography") or {}).get(group)) or {}
    return cfg.get("color") or default


def _typo_style(s, group, default_color, size_override=None):
    """Inline CSS string for a text group, from global typography settings (with defaults)."""
    d = dict(_TYPO_DEFAULTS.get(group, _TYPO_DEFAULTS["contact"]))
    cfg = ((s.get("typography") or {}).get(group)) or {}
    for k, v in cfg.items():
        if v is not None and v != "":
            d[k] = v
    fam = _FONT_STACKS.get(d.get("font") or "Arial", _FONT_STACKS["Arial"])
    col = d.get("color") or default_color
    size = int(size_override or d.get("size") or 13)
    css = f"font-family:{fam};font-size:{size}px;color:{col};font-weight:{'700' if d.get('bold') else '400'};"
    if d.get("italic"):
        css += "font-style:italic;"
    if d.get("underline"):
        css += "text-decoration:underline;"
    return css


def build_signature_html(user, s, include_style=True):
    color = s.get("primary_color") or "#2563EB"
    logo_w = int(s.get("logo_width") or 86)
    banner_w = int(s.get("banner_width") or 600)
    name = _esc(user.get("name") or "Nom Prénom")
    title = _esc(user.get("title") or "")
    dept = _esc(user.get("department") or "")
    email = _esc(user.get("email") or "")
    ext = _esc(user.get("phone_ext") or "")
    direct = _esc(user.get("direct_line") or "")
    phone_main = _esc(s.get("phone_main") or "")
    company = _esc(s.get("company_name") or "")
    website = s.get("website") or ""
    address = _esc(s.get("address") or "")
    disclaimer = _clean_html(s.get("disclaimer") or "")
    avatar = user.get("avatar_url") or ""
    company_logo = s.get("logo_url") or ""
    logo_pos = s.get("logo_position") or "left"
    is_modern = (_resolve_banner(user, s).get("layout") or "classic") == "modern"
    left_img = avatar or (company_logo if (logo_pos == "left" and not is_modern) else "")
    show_logo_above_name = bool(company_logo) and (logo_pos == "above_name" or (logo_pos == "left" and (is_modern or bool(avatar))))
    show_logo_after_address = bool(company_logo) and logo_pos == "after_address"
    show_logo_before_banner = bool(company_logo) and logo_pos == "before_banner"
    show_logo_after_banner = bool(company_logo) and logo_pos == "after_banner"
    top_logo_w = min(logo_w + 30, 140)
    company_logo_tag = (f'<img src="{_esc(company_logo)}" width="{logo_w}" style="display:block;width:{logo_w}px;max-width:100%;border-radius:4px;" alt="{company}" />') if company_logo else ""
    avatar_w = int(user.get("avatar_width") or 0) or logo_w
    left_w = avatar_w if avatar else logo_w
    booking_url = _norm_url(user.get("booking_url") or "")

    title_raw_lines = [l.strip() for l in (user.get("title") or "").split("\n") if l.strip()]
    title_parts = [_esc(l) for l in title_raw_lines] + ([dept] if dept else [])
    title_line = "<br>".join(title_parts)
    phone_parts = [p for p in [phone_main, (f"poste {ext}" if ext else "")] if p]
    phone_str = " ".join(phone_parts)

    cstyle = _typo_style(s, "contact", "#1f2937")
    ccolor = _typo_color(s, "contact", "#1f2937")
    astyle = _typo_style(s, "address", "#6b7280")
    tel_r = _typo_style_contact(s, "tel", "#1f2937")
    dir_r = _typo_style_contact(s, "direct", "#1f2937")
    eml_r = _typo_style_contact(s, "courriel", "#1f2937")
    web_r = _typo_style_contact(s, "website", "#1f2937")

    lines = []
    if phone_str:
        lines.append(f'<tr><td class="sf-text" style="padding:1px 0;{tel_r["css"]}"><span class="sf-accent" style="color:{tel_r["label_color"] or color};font-weight:700;">Tél&nbsp;</span>{phone_str}</td></tr>')
    if direct:
        lines.append(f'<tr><td class="sf-text" style="padding:1px 0;{dir_r["css"]}"><span class="sf-accent" style="color:{dir_r["label_color"] or color};font-weight:700;">Ligne directe&nbsp;</span>{direct}</td></tr>')
    if email:
        lines.append(f'<tr><td class="sf-text" style="padding:1px 0;{eml_r["css"]}"><span class="sf-accent" style="color:{eml_r["label_color"] or color};font-weight:700;">Courriel&nbsp;</span><a href="mailto:{email}" style="color:{eml_r["color"]};text-decoration:none;">{email}</a></td></tr>')
    if website:
        lines.append(f'<tr><td class="sf-text" style="padding:1px 0;{web_r["css"]}"><span class="sf-accent" style="color:{web_r["label_color"] or color};font-weight:700;">Web&nbsp;</span><a href="{_esc(_norm_url(website))}" style="color:{web_r["color"]};text-decoration:none;">{_esc(website)}</a></td></tr>')
    if address:
        lines.append(f'<tr><td class="sf-muted" style="padding:1px 0;{astyle}">{address}</td></tr>')

    social = s.get("social", {}) or {}
    icons_mode = (s.get("social_style") or "icons") == "icons"
    social_links = []
    for k, (label, c) in _SOCIAL.items():
        if social.get(k):
            href = _esc(_norm_url(social[k]))
            if icons_mode:
                icon = f"{BACKEND_PUBLIC_URL}/api/social-icons/{k}.png"
                social_links.append(f'<a href="{href}" style="text-decoration:none;display:inline-block;" target="_blank"><img src="{icon}" width="24" height="24" style="display:inline-block;border:0;width:24px;height:24px;vertical-align:middle;border-radius:5px;" alt="{label}" /></a>')
            else:
                social_links.append(f'<a href="{href}" class="sf-social" style="color:{c};text-decoration:none;font-weight:600;font-size:12px;">{label}</a>')
    sep = "&nbsp;&nbsp;" if icons_mode else ' <span style="color:#d1d5db;">|</span> '
    social_row = f'<tr><td style="padding-top:6px;font-family:Arial,Helvetica,sans-serif;">{sep.join(social_links)}</td></tr>' if social_links else ""

    logo_cell = (f'<td class="sf-bar" style="vertical-align:top;padding-right:18px;border-right:3px solid {color};">'
                 f'<img src="{_esc(left_img)}" width="{left_w}" style="display:block;width:{left_w}px;border-radius:6px;" alt="{company}" /></td>') if left_img else ""
    name_bar = bool(s.get("name_top_border"))
    name_bar_color = s.get("name_top_border_color") or "#111827"
    name_bar_w = int(s.get("name_top_border_width") or 2)
    name_bar_row = (f'<tr><td style="padding-bottom:6px;"><div style="border-top:{name_bar_w}px solid {name_bar_color};font-size:1px;line-height:1px;height:1px;">&nbsp;</div></td></tr>') if name_bar else ""
    name_bar_div = (f'<div style="border-top:{name_bar_w}px solid {name_bar_color};font-size:1px;line-height:1px;height:1px;margin-bottom:6px;">&nbsp;</div>') if name_bar else ""
    logo_above_name_row = (f'<tr><td style="padding-bottom:8px;">{company_logo_tag}</td></tr>') if show_logo_above_name else ""
    logo_after_address_row = (f'<tr><td style="padding-top:8px;padding-bottom:2px;">{company_logo_tag}</td></tr>') if show_logo_after_address else ""
    identity_border = "" if left_img else f'border-left:3px solid {color};padding-left:16px;'
    identity = (
        f'<td style="vertical-align:top;padding-left:{"18px" if left_img else "0"};{identity_border}">'
        f'<table cellpadding="0" cellspacing="0" border="0" style="border-collapse:collapse;">'
        f'{logo_above_name_row}'
        f'{name_bar_row}'
        f'<tr><td class="sf-name" style="{_typo_style(s, "name", "#0f172a")}padding-bottom:2px;">{name}</td></tr>'
        + (f'<tr><td class="sf-accent" style="{_typo_style(s, "title", color)}padding-bottom:2px;">{title_line}</td></tr>' if title_line else "")
        + (f'<tr><td class="sf-name" style="font-family:Arial,Helvetica,sans-serif;font-size:13px;font-weight:600;color:#0f172a;padding-bottom:6px;">{company}</td></tr>' if company else "")
        + "".join(lines)
        + logo_after_address_row
        + social_row
        + "</table></td>"
    )
    rows = []
    banner = _resolve_banner(user, s)
    layout = banner.get("layout") or "classic"

    if layout == "modern":
        parts = []
        if phone_str:
            parts.append(f'<span style="{tel_r["css"]}"><span class="sf-accent" style="color:{tel_r["label_color"] or color};font-weight:700;">Tél</span> {phone_str}</span>')
        if direct:
            parts.append(f'<span style="{dir_r["css"]}"><span class="sf-accent" style="color:{dir_r["label_color"] or color};font-weight:700;">Direct</span> {direct}</span>')
        if email:
            parts.append(f'<a href="mailto:{email}" style="{eml_r["css"]}text-decoration:none;">{email}</a>')
        if website:
            parts.append(f'<a href="{_esc(_norm_url(website))}" style="{web_r["css"]}text-decoration:none;">{_esc(website)}</a>')
        contact_inline = ' &nbsp;<span style="color:#d1d5db;">·</span>&nbsp; '.join(parts)
        photo_cell = (f'<td style="vertical-align:top;padding-right:16px;"><img src="{_esc(avatar)}" width="{avatar_w}" style="display:block;width:{avatar_w}px;border-radius:6px;" alt="Photo" /></td>') if avatar else ""
        content_logo = (f'<div style="margin-bottom:8px;">{company_logo_tag}</div>') if show_logo_above_name else ""
        logo_after_address_div = (f'<div style="padding-top:8px;">{company_logo_tag}</div>') if show_logo_after_address else ""
        content_colspan = "" if avatar else ' colspan="2"'
        modern = (
            photo_cell
            + f'<td{content_colspan} class="sf-bar-left" style="border-left:4px solid {color};padding:2px 0 2px 16px;">'
            f'{content_logo}'
            f'{name_bar_div}'
            f'<div class="sf-name" style="{_typo_style(s, "name", "#0f172a", size_override=(s.get("typography") or {}).get("name", {}).get("size"))}">{name}</div>'
            + (f'<div class="sf-accent" style="{_typo_style(s, "title", color)}padding-top:1px;">{title_line}</div>' if title_line else "")
            + (f'<div class="sf-name" style="font-family:Arial,Helvetica,sans-serif;font-size:13px;font-weight:600;color:#0f172a;padding-top:1px;">{company}</div>' if company else "")
            + (f'<div class="sf-text" style="{cstyle}padding-top:6px;">{contact_inline}</div>' if contact_inline else "")
            + (f'<div class="sf-muted" style="{astyle}padding-top:2px;">{address}</div>' if address else "")
            + logo_after_address_div
            + (f'<div style="padding-top:6px;">{sep.join(social_links)}</div>' if social_links else "")
            + '</td>'
        )
        rows.append(f"<tr>{modern}</tr>")
    else:
        rows.append(f"<tr>{logo_cell}{identity}</tr>")

    if show_logo_before_banner:
        rows.append(f'<tr><td colspan="2" style="padding-top:12px;">{company_logo_tag}</td></tr>')

    if banner["gif_url"]:
        img = f'<img src="{_esc(banner["gif_url"])}" width="{banner_w}" style="display:block;width:{banner_w}px;max-width:100%;border-radius:8px;border:0;" alt="Bannière" />'
        link = _norm_url(banner["banner_link"])
        if link and user.get("id"):
            track = f'{BACKEND_PUBLIC_URL}/api/track/click?u={_esc(user["id"])}&b={urllib.parse.quote(banner["key"] or "Défaut")}&url={urllib.parse.quote(link)}'
            img = f'<a href="{_esc(track)}" target="_blank" style="text-decoration:none;">{img}</a>'
        elif link:
            img = f'<a href="{_esc(link)}" target="_blank" style="text-decoration:none;">{img}</a>'
        rows.append(f'<tr><td colspan="2" style="padding-top:16px;">{img}</td></tr>')

    if show_logo_after_banner:
        rows.append(f'<tr><td colspan="2" style="padding-top:12px;">{company_logo_tag}</td></tr>')

    if booking_url:
        rows.append(f'<tr><td colspan="2" style="padding-top:12px;"><a href="{_esc(booking_url)}" target="_blank" style="display:inline-block;background:{color};color:#ffffff;text-decoration:none;font-family:Arial,Helvetica,sans-serif;font-size:12px;font-weight:700;padding:9px 16px;border-radius:6px;">Réservez une heure pour me rencontrer</a></td></tr>')

    if disclaimer:
        rows.append(f'<tr><td colspan="2" style="padding-top:14px;"><div class="sf-disc sf-disc-line" style="border-top:1px solid #e5e7eb;padding-top:8px;font-family:Arial,Helvetica,sans-serif;font-size:10px;line-height:1.4;color:#9ca3af;max-width:600px;">{disclaimer}</div></td></tr>')

    style_block = _dark_style_block(_dark_accent_for(color)) if include_style else ""
    html = f'{style_block}<table cellpadding="0" cellspacing="0" border="0" style="border-collapse:collapse;font-family:Arial,Helvetica,sans-serif;">{"".join(rows)}</table>'
    return _rewrite_asset_urls(html)


def _signature_file(sig_html):
    return f'<!DOCTYPE html><html><head><meta charset="utf-8"></head><body>{sig_html}</body></html>'


# ---------------------------------------------------------------------------
# Department banners
# ---------------------------------------------------------------------------
@api_router.post("/departments")
async def create_department(data: DeptCreate, admin: dict = Depends(require_perm("gif", "employees", "setup"))):
    settings = await load_settings()
    banners = settings.get("department_banners", [])
    if any(b["name"].strip().lower() == data.name.strip().lower() for b in banners):
        raise HTTPException(status_code=400, detail="Ce département existe déjà")
    banner = DeptBanner(id=str(uuid.uuid4()), name=data.name.strip()).model_dump()
    banners.append(banner)
    await db.settings.update_one({"key": "global"}, {"$set": {"department_banners": banners}}, upsert=True)
    return banner


@api_router.delete("/departments/{dept_id}")
async def delete_department(dept_id: str, admin: dict = Depends(require_perm("gif", "employees", "setup"))):
    settings = await load_settings()
    banners = [b for b in settings.get("department_banners", []) if b["id"] != dept_id]
    await db.settings.update_one({"key": "global"}, {"$set": {"department_banners": banners}})
    return {"ok": True}


# ---------------------------------------------------------------------------
# SMTP configuration + email sending
# ---------------------------------------------------------------------------
@api_router.get("/smtp")
async def get_smtp(admin: dict = Depends(require_perm("email"))):
    cfg = await db.smtp_config.find_one({"key": "global"}) or {}
    return {
        "host": cfg.get("host", ""), "port": cfg.get("port", 587),
        "username": cfg.get("username", ""), "from_email": cfg.get("from_email", ""),
        "from_name": cfg.get("from_name", ""), "has_password": bool(cfg.get("password")),
    }


@api_router.put("/smtp")
async def put_smtp(data: SmtpConfigInput, admin: dict = Depends(require_perm("email"))):
    existing = await db.smtp_config.find_one({"key": "global"}) or {}
    payload = {"key": "global", "host": data.host.strip(), "port": int(data.port),
               "username": data.username.strip(), "from_email": data.from_email.strip(),
               "from_name": data.from_name.strip()}
    if data.password:
        payload["password"] = data.password
    elif existing.get("password"):
        payload["password"] = existing["password"]
    await db.smtp_config.update_one({"key": "global"}, {"$set": payload}, upsert=True)
    return {"ok": True}


async def _send_email(to, subject, html_body, attachment_html=None, attachment_name="signature.html"):
    cfg = await db.smtp_config.find_one({"key": "global"})
    if not cfg or not cfg.get("host"):
        raise HTTPException(status_code=400, detail="SMTP non configuré. Renseignez les paramètres d'envoi.")
    msg = EmailMessage()
    sender = cfg.get("from_email") or cfg.get("username")
    msg["From"] = f'{cfg.get("from_name", "")} <{sender}>'.strip()
    msg["To"] = to
    msg["Subject"] = subject
    msg.set_content("Votre client de messagerie ne supporte pas le HTML. Voir la pièce jointe.")
    msg.add_alternative(html_body, subtype="html")
    if attachment_html:
        msg.add_attachment(attachment_html.encode("utf-8"), maintype="text",
                           subtype="html", filename=attachment_name)
    port = int(cfg.get("port", 587))
    kwargs = {"use_tls": True} if port == 465 else {"start_tls": True}
    try:
        await aiosmtplib.send(msg, hostname=cfg["host"], port=port,
                              username=cfg.get("username") or None,
                              password=cfg.get("password") or None, timeout=30, **kwargs)
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"SMTP send failed: {e}")
        raise HTTPException(status_code=502, detail=f"Échec de l'envoi SMTP : {e}")


def _email_body(user, signature_html):
    name = user.get("name", "")
    link = f"{BACKEND_PUBLIC_URL}/app/signature"
    return f"""
<div style="font-family:Arial,Helvetica,sans-serif;max-width:660px;margin:0 auto;color:#1f2937;">
  <h2 style="color:#111827;">Bonjour {name},</h2>
  <p>Voici votre <b>signature courriel officielle</b>. Vous pouvez la copier directement dans Outlook,
  ou utiliser le fichier <b>signature.html</b> joint à ce message.</p>
  <div style="background:#f8fafc;border:1px solid #e2e8f0;border-radius:10px;padding:20px;margin:18px 0;">{signature_html}</div>
  <p style="font-size:14px;"><b>Installation dans Outlook :</b> Fichier → Options → Courrier → Signatures →
  créez une signature, cliquez dans la zone d'édition puis collez (Ctrl+V).</p>
  <p style="font-size:14px;">Mettez à jour vos coordonnées ici : <a href="{link}" style="color:#2563EB;">{link}</a></p>
  <p style="font-size:12px;color:#9ca3af;">Message envoyé automatiquement par SigChamp Office.</p>
</div>"""


@api_router.post("/email/test")
async def email_test(data: EmailTestInput, admin: dict = Depends(require_perm("email"))):
    await _send_email(str(data.to), "Test SigChamp — configuration SMTP",
                      "<p>Ceci est un courriel de test. Votre configuration SMTP fonctionne. ✅</p>")
    return {"ok": True}


async def _log_send(user: dict, status: str, error: str = ""):
    now = datetime.now(timezone.utc).isoformat()
    await db.send_log.insert_one({
        "user_id": user.get("id"), "email": user.get("email"), "name": user.get("name"),
        "status": status, "error": error, "ts": now,
    })
    await db.users.update_one({"_id": ObjectId(user["id"])},
                              {"$set": {"last_sent_at": now, "last_send_status": status}})


async def _send_signature_to(user: dict, settings: dict):
    sig = build_signature_html(user, settings)
    try:
        await _send_email(user["email"], "Votre signature courriel officielle",
                          _email_body(user, sig), attachment_html=_signature_file(sig))
        await _log_send(user, "sent")
        return True, ""
    except HTTPException as e:
        await _log_send(user, "failed", str(e.detail))
        raise
    except Exception as e:
        await _log_send(user, "failed", str(e))
        raise


@api_router.post("/email/send/{emp_id}")
async def email_send_one(emp_id: str, admin: dict = Depends(require_perm("email"))):
    user = await db.users.find_one({"_id": _oid(emp_id)})
    if not user:
        raise HTTPException(status_code=404, detail="Employé introuvable")
    user["id"] = str(user["_id"])
    settings = await load_settings()
    await _send_signature_to(user, settings)
    return {"ok": True, "sent_to": user["email"]}


@api_router.post("/email/send-all")
async def email_send_all(admin: dict = Depends(require_perm("email"))):
    settings = await load_settings()
    users = await db.users.find().to_list(1000)
    sent, failed = [], []
    for u in users:
        u["id"] = str(u["_id"])
        try:
            await _send_signature_to(u, settings)
            sent.append(u["email"])
        except Exception as e:
            failed.append({"email": u["email"], "error": str(getattr(e, "detail", e))})
    return {"sent": sent, "failed": failed, "sent_count": len(sent)}


@api_router.post("/email/send-reminders")
async def email_send_reminders(admin: dict = Depends(require_perm("email"))):
    settings = await load_settings()
    users = await db.users.find({"role": "employee", "signature_installed": {"$ne": True}}).to_list(1000)
    sent, failed = [], []
    for u in users:
        u["id"] = str(u["_id"])
        try:
            await _send_signature_to(u, settings)
            sent.append(u["email"])
        except Exception as e:
            failed.append({"email": u["email"], "error": str(getattr(e, "detail", e))})
    return {"sent": sent, "failed": failed, "sent_count": len(sent), "targeted": len(users)}


@api_router.put("/employees/{emp_id}/installed")
async def mark_installed(emp_id: str, admin: dict = Depends(require_perm("employees", "email", "deploy"))):
    user = await db.users.find_one({"_id": _oid(emp_id)})
    if not user:
        raise HTTPException(status_code=404, detail="Employé introuvable")
    new_val = not bool(user.get("signature_installed", False))
    update = {"signature_installed": new_val,
              "installed_at": datetime.now(timezone.utc).isoformat() if new_val else None}
    await db.users.update_one({"_id": _oid(emp_id)}, {"$set": update})
    fresh = await db.users.find_one({"_id": _oid(emp_id)})
    fresh["id"] = str(fresh["_id"])
    return user_to_public(fresh)


@api_router.get("/send-log")
async def get_send_log(admin: dict = Depends(require_perm("email", "analytics"))):
    entries = await db.send_log.find().sort("ts", -1).to_list(200)
    for e in entries:
        e["id"] = str(e.pop("_id"))
    return entries


# ---------------------------------------------------------------------------
# Click tracking + analytics
# ---------------------------------------------------------------------------
@app.get("/api/track/click")
async def track_click(u: str = Query(""), b: str = Query("Défaut"), url: str = Query("")):
    target = url if url.lower().startswith(("http://", "https://")) else "https://www.google.com"
    try:
        await db.clicks.insert_one({"user_id": u, "banner": b, "target": target,
                                    "ts": datetime.now(timezone.utc).isoformat()})
    except Exception as e:
        logger.error(f"click log failed: {e}")
    return RedirectResponse(url=target, status_code=302)


@api_router.get("/analytics/clicks")
async def analytics_clicks(admin: dict = Depends(require_perm("analytics"))):
    clicks = await db.clicks.find().to_list(20000)
    by_banner, by_user = {}, {}
    for c in clicks:
        key = c.get("banner", "Défaut")
        by_banner[key] = by_banner.get(key, 0) + 1
        uid = c.get("user_id", "")
        if uid:
            by_user[uid] = by_user.get(uid, 0) + 1
    users = await db.users.find().to_list(1000)
    umap = {str(x["_id"]): x.get("name", x.get("email", "")) for x in users}
    return {
        "total": len(clicks),
        "by_banner": sorted([{"banner": k, "count": v} for k, v in by_banner.items()], key=lambda x: -x["count"]),
        "by_employee": sorted([{"user_id": k, "name": umap.get(k, "Inconnu"), "count": v}
                               for k, v in by_user.items()], key=lambda x: -x["count"]),
    }


# ---------------------------------------------------------------------------
# Employees (admin)
# ---------------------------------------------------------------------------
@api_router.get("/employees")
async def list_employees(admin: dict = Depends(require_any_management)):
    users = await db.users.find().sort("name", 1).to_list(1000)
    return [user_to_public({**u, "id": str(u["_id"])}) for u in users]


@api_router.post("/employees")
async def create_employee(data: EmployeeCreate, admin: dict = Depends(require_perm("employees"))):
    email = data.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Cet email existe déjà")
    doc = {
        "email": email, "password_hash": hash_password(data.password), "name": data.name,
        "role": "employee", "title": data.title, "phone_ext": data.phone_ext,
        "direct_line": data.direct_line, "department": data.department, "avatar_url": "",
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    res = await db.users.insert_one(doc)
    doc["id"] = str(res.inserted_id)
    return user_to_public(doc)


_IMPORT_MAP = {
    "name": "name", "nom": "name", "full name": "name", "nom complet": "name",
    "email": "email", "courriel": "email", "e-mail": "email", "mail": "email",
    "password": "password", "mot de passe": "password", "motdepasse": "password",
    "title": "title", "poste": "title", "titre": "title", "fonction": "title",
    "department": "department", "departement": "department", "département": "department", "service": "department",
    "phone_ext": "phone_ext", "poste telephonique": "phone_ext", "poste téléphonique": "phone_ext",
    "extension": "phone_ext", "poste tel": "phone_ext",
    "direct_line": "direct_line", "ligne directe": "direct_line", "telephone": "direct_line", "téléphone": "direct_line",
}


def _norm_row(raw: dict) -> dict:
    out = {}
    for k, v in raw.items():
        if k is None:
            continue
        key = _IMPORT_MAP.get(str(k).strip().lower())
        if key:
            out[key] = ("" if v is None else str(v)).strip()
    return out


def _parse_import(filename: str, data: bytes) -> list:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    rows = []
    if ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        header = None
        for r in ws.iter_rows(values_only=True):
            if header is None:
                header = [str(c).strip() if c is not None else "" for c in r]
                continue
            if all(c is None or str(c).strip() == "" for c in r):
                continue
            rows.append(_norm_row(dict(zip(header, r))))
    else:
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        for r in reader:
            rows.append(_norm_row(r))
    return rows


@api_router.post("/employees/import")
async def import_employees(file: UploadFile = File(...), admin: dict = Depends(require_perm("employees"))):
    data = await file.read()
    try:
        rows = _parse_import(file.filename or "import.csv", data)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Fichier illisible : {e}")

    created, skipped, errors = [], [], []
    for i, row in enumerate(rows, start=2):
        email = (row.get("email") or "").lower().strip()
        name = (row.get("name") or "").strip()
        if not email:
            errors.append({"row": i, "error": "Courriel manquant"})
            continue
        if "@" not in email:
            errors.append({"row": i, "error": f"Courriel invalide : {email}"})
            continue
        if await db.users.find_one({"email": email}):
            skipped.append(email)
            continue
        password = row.get("password") or secrets.token_urlsafe(8)
        doc = {
            "email": email, "password_hash": hash_password(password),
            "name": name or email.split("@")[0], "role": "employee",
            "title": row.get("title", ""), "phone_ext": row.get("phone_ext", ""),
            "direct_line": row.get("direct_line", ""), "department": row.get("department", ""),
            "avatar_url": "", "created_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.users.insert_one(doc)
        created.append({"email": email, "name": doc["name"], "password": password})

    return {"created": created, "skipped": skipped, "errors": errors,
            "created_count": len(created), "skipped_count": len(skipped)}


@api_router.put("/employees/{emp_id}")
async def update_employee(emp_id: str, data: ProfileUpdate, admin: dict = Depends(require_perm("employees"))):
    update = {k: v for k, v in data.model_dump().items() if v is not None}
    if update:
        await db.users.update_one({"_id": _oid(emp_id)}, {"$set": update})
    fresh = await db.users.find_one({"_id": _oid(emp_id)})
    if not fresh:
        raise HTTPException(status_code=404, detail="Employé introuvable")
    fresh["id"] = str(fresh["_id"])
    return user_to_public(fresh)


@api_router.delete("/employees/{emp_id}")
async def delete_employee(emp_id: str, admin: dict = Depends(require_perm("employees"))):
    target = await db.users.find_one({"_id": _oid(emp_id)})
    if not target:
        raise HTTPException(status_code=404, detail="Employé introuvable")
    if target.get("role") == "admin":
        raise HTTPException(status_code=400, detail="Impossible de supprimer un administrateur")
    await db.users.delete_one({"_id": _oid(emp_id)})
    return {"ok": True}


class AccessUpdate(BaseModel):
    role: str
    permissions: List[str] = []


@api_router.put("/employees/{emp_id}/access")
async def update_employee_access(emp_id: str, data: AccessUpdate, admin: dict = Depends(require_admin)):
    """Définit le rôle et les permissions de section d'un utilisateur (admin uniquement)."""
    role = "admin" if data.role == "admin" else "employee"
    perms = [] if role == "admin" else [p for p in data.permissions if p in SECTION_PERMS]
    if emp_id == admin["id"] and role != "admin":
        raise HTTPException(status_code=400, detail="Vous ne pouvez pas retirer votre propre rôle administrateur.")
    target = await db.users.find_one({"_id": _oid(emp_id)})
    if not target:
        raise HTTPException(status_code=404, detail="Utilisateur introuvable")
    await db.users.update_one({"_id": _oid(emp_id)}, {"$set": {"role": role, "permissions": perms}})
    fresh = await db.users.find_one({"_id": _oid(emp_id)})
    fresh["id"] = str(fresh["_id"])
    return user_to_public(fresh)


# ---------------------------------------------------------------------------
# Microsoft 365 deployment scripts (PowerShell, pre-filled)
# ---------------------------------------------------------------------------
def _ps_signatures_block(users, settings):
    lines = ["$signatures = @{}"]
    for u in users:
        u["id"] = str(u["_id"])
        sig = build_signature_html(u, settings).replace("\r", "").replace("\n", " ")
        email = u.get("email", "").replace("'", "''")
        lines.append(f"$signatures['{email}'] = @'\n{sig}\n'@")
    return "\n".join(lines)


def _ps_response(script: str, filename: str):
    return Response(content=script, media_type="text/plain; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


DEPLOY_API_TOKEN = os.environ.get("DEPLOY_API_TOKEN", "")


async def get_deploy_token() -> str:
    doc = await db.config.find_one({"key": "deploy"})
    if doc and doc.get("token"):
        return doc["token"]
    return DEPLOY_API_TOKEN


def _deploy_urls(tok: str) -> dict:
    base = BACKEND_PUBLIC_URL.rstrip("/")
    return {
        "has_token": bool(tok),
        "exchange_url": f"{base}/api/deploy/exchange-script?token={tok}" if tok else "",
        "gpo_url": f"{base}/api/deploy/gpo-script?token={tok}" if tok else "",
    }


async def deploy_auth(request: Request, token: Optional[str] = Query(default=None),
                      creds: Optional[HTTPAuthorizationCredentials] = Depends(security)):
    """Allow either an admin JWT (interactive UI) or a static token (unattended scheduled tasks)."""
    current = await get_deploy_token()
    if current and token and token == current:
        return {"role": "admin", "via": "token"}
    user = await get_current_user(request, creds)
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Accès réservé aux administrateurs")
    return user


@api_router.get("/deploy/info")
async def deploy_info(admin: dict = Depends(require_perm("deploy", "install", "m365"))):
    doc = await db.config.find_one({"key": "deploy"}) or {}
    res = _deploy_urls(await get_deploy_token())
    res["rotated_at"] = doc.get("rotated_at")
    return res


@api_router.post("/deploy/rotate-token")
async def rotate_deploy_token(admin: dict = Depends(require_perm("deploy", "install"))):
    new_token = secrets.token_urlsafe(24)
    now = datetime.now(timezone.utc).isoformat()
    await db.config.update_one({"key": "deploy"},
                               {"$set": {"token": new_token, "rotated_at": now}}, upsert=True)
    res = _deploy_urls(new_token)
    res["rotated_at"] = now
    return res


@api_router.get("/deploy/exchange-script")
async def deploy_exchange_script(_auth: dict = Depends(deploy_auth)):
    settings = await load_settings()
    users = await db.users.find({"role": "employee"}).to_list(1000)
    block = _ps_signatures_block(users, settings)
    script = f"""# =====================================================================
# SigChamp Office - Signatures Exchange Online (une regle par employe)
# ---------------------------------------------------------------------
# Ajoute automatiquement la signature en bas de chaque courriel sortant.
# Aucun copier-coller cote employe.
#
# ETAPES (a executer par l'admin Microsoft 365) :
#   1. Install-Module ExchangeOnlineManagement -Scope CurrentUser
#   2. Connect-ExchangeOnline -UserPrincipalName admin@votredomaine.com
#   3. Executez ce script :  .\\sigchamp-exchange-signatures.ps1
# =====================================================================

{block}

foreach ($email in $signatures.Keys) {{
    $ruleName = "SigChamp - $email"
    $html = $signatures[$email]
    if (Get-TransportRule -Identity $ruleName -ErrorAction SilentlyContinue) {{
        Set-TransportRule -Identity $ruleName -From $email `
            -ApplyHtmlDisclaimerLocation Append -ApplyHtmlDisclaimerText $html `
            -ApplyHtmlDisclaimerFallbackAction Wrap
        Write-Host "Mise a jour : $ruleName"
    }} else {{
        New-TransportRule -Name $ruleName -From $email `
            -ApplyHtmlDisclaimerLocation Append -ApplyHtmlDisclaimerText $html `
            -ApplyHtmlDisclaimerFallbackAction Wrap
        Write-Host "Cree : $ruleName"
    }}
}}
Write-Host "Termine. $($signatures.Count) signature(s) configuree(s)."
"""
    return _ps_response(script, "sigchamp-exchange-signatures.ps1")


@api_router.get("/deploy/gpo-script")
async def deploy_gpo_script(_auth: dict = Depends(deploy_auth)):
    settings = await load_settings()
    users = await db.users.find({"role": "employee"}).to_list(1000)
    block = _ps_signatures_block(users, settings)
    script = f"""# =====================================================================
# SigChamp Office - Deploiement signature Outlook via GPO
# ---------------------------------------------------------------------
# A deployer en script d'ouverture de session :
#   Configuration utilisateur > Strategies > Parametres Windows > Scripts
#   > Ouverture de session > ajouter ce .ps1
# Definit la signature par defaut d'Outlook (bureau) pour chaque employe.
# =====================================================================

{block}

# Detecte le courriel de l'utilisateur courant (Active Directory)
$email = $null
try {{
    $email = ([ADSISearcher]"(&(objectCategory=User)(sAMAccountName=$env:USERNAME))").FindOne().Properties["mail"][0]
}} catch {{}}

if (-not $email) {{ Write-Host "Courriel introuvable pour $env:USERNAME"; exit }}

if ($signatures.ContainsKey($email)) {{
    $sigName = "Signature entreprise"
    $sigDir = Join-Path $env:APPDATA "Microsoft\\Signatures"
    New-Item -ItemType Directory -Force -Path $sigDir | Out-Null
    $html = $signatures[$email]
    [System.IO.File]::WriteAllText((Join-Path $sigDir "$sigName.htm"), $html, [System.Text.Encoding]::UTF8)
    [System.IO.File]::WriteAllText((Join-Path $sigDir "$sigName.txt"), "", [System.Text.Encoding]::UTF8)

    # Signature par defaut (Outlook 2016 / 2019 / 365 = Office 16.0)
    $base = "HKCU:\\Software\\Microsoft\\Office\\16.0\\Common\\MailSettings"
    New-Item -Path $base -Force | Out-Null
    Set-ItemProperty -Path $base -Name "NewSignature" -Value $sigName
    Set-ItemProperty -Path $base -Name "ReplySignature" -Value $sigName
    Write-Host "Signature installee pour $email"
}} else {{
    Write-Host "Aucune signature SigChamp trouvee pour $email"
}}
"""
    return _ps_response(script, "sigchamp-outlook-gpo.ps1")


# ---------------------------------------------------------------------------
# Microsoft 365 direct connection (Graph list users + Exchange push)
# ---------------------------------------------------------------------------
class M365ConfigInput(BaseModel):
    tenant_id: str = ""
    client_id: str = ""
    client_secret: Optional[str] = None
    tenant_domain: str = ""


class M365PushInput(BaseModel):
    emails: List[str]
    fallback: str = "Ignore"


async def _m365_cfg():
    return await db.config.find_one({"key": "m365"}) or {}


def _m365_token(cfg: dict, scope: str) -> str:
    r = requests.post(
        f"https://login.microsoftonline.com/{cfg['tenant_id']}/oauth2/v2.0/token",
        data={"client_id": cfg["client_id"], "client_secret": cfg.get("client_secret", ""),
              "scope": scope, "grant_type": "client_credentials"}, timeout=30)
    if r.status_code != 200:
        raise HTTPException(status_code=502, detail=f"Auth M365 échouée : {r.text[:300]}")
    return r.json()["access_token"]


def _graph_list_users(cfg: dict) -> list:
    tok = _m365_token(cfg, "https://graph.microsoft.com/.default")
    url = "https://graph.microsoft.com/v1.0/users?$select=id,displayName,userPrincipalName,mail,accountEnabled&$top=999"
    out = []
    seen = set()
    while url:
        r = requests.get(url, headers={"Authorization": f"Bearer {tok}"}, timeout=30)
        if r.status_code != 200:
            raise HTTPException(status_code=502, detail=f"Graph : {r.text[:300]}")
        body = r.json()
        for u in body.get("value", []):
            em = (u.get("mail") or u.get("userPrincipalName") or "").strip()
            key = em.lower()
            if not key or key in seen:
                continue
            seen.add(key)
            out.append({"id": u["id"], "name": u.get("displayName") or "",
                        "email": em, "enabled": u.get("accountEnabled", True)})
        url = body.get("@odata.nextLink")
    return out





def _clean_ext(phones):
    if not phones:
        return ""
    v = phones[0] if isinstance(phones, list) else str(phones)
    return v.strip().lstrip("xX").strip()


def _graph_get_user(cfg: dict, email: str) -> dict:
    """Récupère nom complet / poste / téléphone depuis Microsoft 365 pour enrichir la signature."""
    try:
        tok = _m365_token(cfg, "https://graph.microsoft.com/.default")
        sel = "displayName,jobTitle,department,officeLocation,businessPhones,mobilePhone"
        r = requests.get(
            f"https://graph.microsoft.com/v1.0/users/{urllib.parse.quote(email)}?$select={sel}",
            headers={"Authorization": f"Bearer {tok}"}, timeout=20)
        if r.status_code == 200:
            return r.json()
    except Exception:
        pass
    return {}


def _m365_get_signature(cfg: dict, email: str) -> str:
    """Récupère la signature que l'employé a définie lui-même dans Outlook (OWA)."""
    try:
        r = _exo_invoke(cfg, "Get-MailboxMessageConfiguration", {"Identity": email})
        if r.status_code < 300:
            v = r.json().get("value", [])
            if v:
                return v[0].get("SignatureHtml") or ""
    except Exception:
        pass
    return ""


def _m365_provision_fields(cfg: dict, email: str) -> dict:
    g = _graph_get_user(cfg, email)
    return {
        "name": g.get("displayName") or email.split("@")[0],
        "title": g.get("jobTitle") or "",
        "department": g.get("department") or "",
        "phone_ext": _clean_ext(g.get("businessPhones")),
        "direct_line": g.get("mobilePhone") or "",
        "m365_signature_html": _m365_get_signature(cfg, email),
    }


def _exo_invoke(cfg: dict, cmdlet: str, params: dict):
    tok = _m365_token(cfg, "https://outlook.office365.com/.default")
    # L'API adminapi Exchange n'accepte PAS un domaine vanité (ex. champagneur.qc.ca) => 401.
    # On cible le tenant par son GUID, qui est toujours accepté.
    tid = cfg["tenant_id"]
    headers = {"Authorization": f"Bearer {tok}", "Content-Type": "application/json",
               "X-AnchorMailbox": f"APP:SystemMailbox{{bb558c35-97f1-4cb9-8ff7-d53741dc928c}}@{tid}"}
    return requests.post(f"https://outlook.office365.com/adminapi/beta/{tid}/InvokeCommand",
                         headers=headers, json={"CmdletInput": {"CmdletName": cmdlet, "Parameters": params}}, timeout=60)


def _exo_error(r):
    """Transforme une réponse d'erreur Exchange (souvent binaire) en message clair et actionnable."""
    code = r.status_code
    base = None
    try:
        j = r.json()
        err = j.get("error")
        if isinstance(err, dict):
            base = err.get("message")
        base = base or j.get("message")
    except Exception:
        base = None
    if not base:
        try:
            txt = "".join(ch for ch in r.text if ch.isprintable()).strip()
            base = txt[:250] if txt else None
        except Exception:
            base = None
    if code == 401:
        return ("Exchange Online a refusé la requête (401 Non autorisé). Le jeton est valide, mais l'application "
                "Azure n'a pas le droit d'exécuter des commandes Exchange. Dans le portail Microsoft Entra (Azure AD) : "
                "1) « API permissions » → ajoutez « Office 365 Exchange Online → Exchange.ManageAsApp » (type Application), "
                "puis cliquez « Grant admin consent » ; 2) « Roles and administrators » → attribuez à l'application le rôle "
                "« Exchange Administrator » (ou « Global Administrator »). Patientez quelques minutes puis réessayez.")
    if code == 403:
        return ("Exchange Online a refusé la requête (403 Interdit). L'application n'a pas les rôles RBAC nécessaires. "
                "Attribuez-lui le rôle « Exchange Administrator » dans Microsoft Entra, puis réessayez.")
    return base or f"Exchange Online a répondu avec le statut {code}."


@api_router.get("/m365/config")
async def get_m365_config(admin: dict = Depends(require_perm("m365", "deploy", "setup"))):
    cfg = await _m365_cfg()
    return {"tenant_id": cfg.get("tenant_id", ""), "client_id": cfg.get("client_id", ""),
            "tenant_domain": cfg.get("tenant_domain", ""), "has_secret": bool(cfg.get("client_secret")),
            "connected": bool(cfg.get("tenant_id") and cfg.get("client_id") and cfg.get("client_secret"))}


@api_router.put("/m365/config")
async def put_m365_config(data: M365ConfigInput, admin: dict = Depends(require_perm("m365", "setup"))):
    existing = await _m365_cfg()
    payload = {"key": "m365", "tenant_id": data.tenant_id.strip(), "client_id": data.client_id.strip(),
               "tenant_domain": data.tenant_domain.strip()}
    if data.client_secret:
        payload["client_secret"] = data.client_secret
    elif existing.get("client_secret"):
        payload["client_secret"] = existing["client_secret"]
    await db.config.update_one({"key": "m365"}, {"$set": payload}, upsert=True)
    return {"ok": True}


@api_router.get("/m365/users")
async def m365_users(admin: dict = Depends(require_perm("m365", "deploy"))):
    cfg = await _m365_cfg()
    if not cfg.get("client_secret"):
        raise HTTPException(status_code=400, detail="Microsoft 365 non configuré.")
    users = await asyncio.to_thread(_graph_list_users, cfg)
    known = {u["email"].lower() async for u in _iter_employee_emails()}
    for u in users:
        u["has_employee"] = u["email"].lower() in known
    return users


async def _iter_employee_emails():
    for u in await db.users.find({}, {"email": 1}).to_list(2000):
        yield {"email": u.get("email", "")}


class RoamingInput(BaseModel):
    postpone: bool


_sync_progress = {"running": False, "total": 0, "done": 0, "synced": 0, "failed": []}


@api_router.get("/m365/roaming")
async def m365_roaming_get(admin: dict = Depends(require_perm("m365", "deploy"))):
    cfg = await _m365_cfg()
    r = await asyncio.to_thread(_exo_invoke, cfg, "Get-OrganizationConfig", {})
    if r.status_code >= 300:
        raise HTTPException(status_code=400, detail=_exo_error(r))
    o = r.json().get("value", [{}])[0]
    return {"postponed": bool(o.get("PostponeRoamingSignaturesUntilLater"))}


@api_router.post("/m365/roaming")
async def m365_roaming_set(data: RoamingInput, admin: dict = Depends(require_perm("m365", "deploy"))):
    cfg = await _m365_cfg()
    r = await asyncio.to_thread(_exo_invoke, cfg, "Set-OrganizationConfig",
                                {"PostponeRoamingSignaturesUntilLater": data.postpone})
    if r.status_code >= 300:
        raise HTTPException(status_code=400, detail=_exo_error(r))
    return {"postponed": data.postpone}


async def _run_sync(cfg: dict, emails: list):
    _sync_progress.update({"running": True, "total": len(emails), "done": 0, "synced": 0, "failed": []})
    for email in emails:
        try:
            fields = await asyncio.to_thread(_m365_provision_fields, cfg, email)
            fields["m365_synced_at"] = datetime.now(timezone.utc).isoformat()
            await db.users.update_one({"email": email}, {"$set": fields})
            _sync_progress["synced"] += 1
        except Exception as e:
            _sync_progress["failed"].append({"email": email, "error": str(e)})
        _sync_progress["done"] += 1
    _sync_progress["running"] = False


@api_router.post("/m365/sync")
async def m365_sync(admin: dict = Depends(require_perm("m365", "deploy"))):
    if _sync_progress["running"]:
        raise HTTPException(status_code=409, detail="Une synchronisation est déjà en cours.")
    cfg = await _m365_cfg()
    if not cfg.get("client_secret"):
        raise HTTPException(status_code=400, detail="Microsoft 365 non configuré.")
    linked = await db.users.find({"m365_linked": True}, {"email": 1}).to_list(5000)
    emails = [(u.get("email") or "").lower() for u in linked if u.get("email")]
    asyncio.create_task(_run_sync(cfg, emails))
    return {"started": True, "total": len(emails)}


@api_router.get("/m365/sync/status")
async def m365_sync_status(admin: dict = Depends(require_perm("m365", "deploy"))):
    return _sync_progress


@api_router.post("/m365/import")
async def m365_import(data: M365PushInput, admin: dict = Depends(require_perm("m365", "deploy"))):
    cfg = await _m365_cfg()
    if not cfg.get("client_secret"):
        raise HTTPException(status_code=400, detail="Microsoft 365 non configuré.")
    imported, updated, failed = [], [], []
    for email in data.emails:
        email = email.strip().lower()
        if not email:
            continue
        try:
            fields = await asyncio.to_thread(_m365_provision_fields, cfg, email)
            doc = {**fields, "email": email, "role": "employee", "m365_linked": True,
                   "m365_synced_at": datetime.now(timezone.utc).isoformat()}
            existing = await db.users.find_one({"email": email})
            if existing:
                await db.users.update_one({"_id": existing["_id"]}, {"$set": doc})
                updated.append(email)
            else:
                doc["created_at"] = datetime.now(timezone.utc).isoformat()
                await db.users.insert_one(doc)
                imported.append(email)
        except Exception as e:
            failed.append({"email": email, "error": str(e)})
    return {"imported": imported, "updated": updated, "failed": failed,
            "imported_count": len(imported), "updated_count": len(updated)}


@api_router.post("/m365/push")
async def m365_push(data: M365PushInput, admin: dict = Depends(require_perm("m365", "deploy"))):
    cfg = await _m365_cfg()
    if not cfg.get("client_secret") or not cfg.get("tenant_domain"):
        raise HTTPException(status_code=400, detail="Microsoft 365 non configuré (domaine + identifiants requis).")
    if data.fallback not in ("Wrap", "Ignore", "Reject"):
        data.fallback = "Ignore"
    settings = await load_settings()
    applied, failed = [], []
    for email in data.emails:
        email = email.strip().lower()
        if not email:
            continue
        emp = await db.users.find_one({"email": email})
        g = await asyncio.to_thread(_graph_get_user, cfg, email)
        emp = emp or {}
        user = {
            **emp,
            "email": email,
            "id": str(emp["_id"]) if emp.get("_id") else None,
            "name": emp.get("name") or g.get("displayName") or email.split("@")[0],
            "title": emp.get("title") or g.get("jobTitle") or "",
            "department": emp.get("department") or g.get("department") or "",
            "phone_ext": emp.get("phone_ext") or _clean_ext(g.get("businessPhones")),
            "direct_line": emp.get("direct_line") or g.get("mobilePhone") or "",
            "booking_url": emp.get("booking_url") or "",
        }
        # Version compacte (sans bloc <style>) : les règles de flux Exchange suppriment le CSS
        # <style> de toute façon et limitent la taille du disclaimer (~5000 caractères).
        # Espace en tête pour que la signature ne colle pas au texte du message (règle de flux = ajout direct).
        html = '<div style="height:18px;line-height:18px;">&nbsp;</div>' + build_signature_html(user, settings, include_style=False)
        rule = f"SigChamp - {email}"
        try:
            # Méthode : règle de flux Exchange (disclaimer serveur). La signature est ajoutée à
            # CHAQUE courriel sortant, quel que soit le client (OWA, Nouveau Outlook, Classique, mobile),
            # sans dépendre des signatures itinérantes (roaming) — que Microsoft n'expose par aucune API.
            fb = data.fallback  # Wrap / Ignore / Reject
            params = {"From": [email],
                      "ApplyHtmlDisclaimerLocation": "Append",
                      "ApplyHtmlDisclaimerText": html,
                      "ApplyHtmlDisclaimerFallbackAction": fb}
            exists = await asyncio.to_thread(_exo_invoke, cfg, "Get-TransportRule", {"Identity": rule})
            if exists.status_code < 300 and (exists.json().get("value") or []):
                r = await asyncio.to_thread(_exo_invoke, cfg, "Set-TransportRule", {"Identity": rule, **params})
            else:
                r = await asyncio.to_thread(_exo_invoke, cfg, "New-TransportRule", {"Name": rule, **params})
            # Éviter tout doublon si la boîte repasse un jour en mode « legacy » : effacer la
            # signature de compte éventuellement posée par une version précédente de SigChamp.
            try:
                await asyncio.to_thread(_exo_invoke, cfg, "Set-MailboxMessageConfiguration",
                                        {"Identity": email, "AutoAddSignature": False,
                                         "AutoAddSignatureOnReply": False, "AutoAddSignatureOnMobile": False,
                                         "DefaultSignature": "", "DefaultSignatureOnReply": "",
                                         "DeleteSignatureName": "SigChamp"})
            except Exception:
                pass
            if r.status_code < 300:
                applied.append(email)
                if emp:
                    await db.users.update_one({"_id": emp["_id"]},
                                              {"$set": {"m365_pushed_at": datetime.now(timezone.utc).isoformat()}})
            else:
                failed.append({"email": email, "error": _exo_error(r)})
        except HTTPException as e:
            failed.append({"email": email, "error": str(e.detail)})
        except Exception as e:
            failed.append({"email": email, "error": str(e)})
    return {"applied": applied, "failed": failed, "applied_count": len(applied)}


@api_router.get("/m365/test")
async def m365_test(admin: dict = Depends(require_perm("m365", "deploy"))):
    cfg = await _m365_cfg()
    if not cfg.get("client_secret") or not cfg.get("client_id") or not cfg.get("tenant_id"):
        raise HTTPException(status_code=400, detail="Renseignez d'abord Tenant ID, Client ID et Secret.")

    def _run():
        result = {"graph_ok": False, "graph_error": None, "exchange_ok": False, "exchange_error": None}
        try:
            tok = _m365_token(cfg, "https://graph.microsoft.com/.default")
            r = requests.get("https://graph.microsoft.com/v1.0/users?$top=1",
                             headers={"Authorization": f"Bearer {tok}"}, timeout=20)
            result["graph_ok"] = r.status_code == 200
            if r.status_code != 200:
                result["graph_error"] = r.text[:200]
        except HTTPException as e:
            result["graph_error"] = str(e.detail)
        except Exception as e:
            result["graph_error"] = str(e)
        try:
            r = _exo_invoke(cfg, "Get-OrganizationConfig", {})
            if r.status_code < 300:
                result["exchange_ok"] = True
            else:
                result["exchange_error"] = _exo_error(r)
        except HTTPException as e:
            result["exchange_error"] = str(e.detail)
        except Exception as e:
            result["exchange_error"] = str(e)
        return result

    return await asyncio.to_thread(_run)


@api_router.post("/m365/remove")
async def m365_remove(data: M365PushInput, admin: dict = Depends(require_perm("m365", "deploy"))):
    cfg = await _m365_cfg()
    if not cfg.get("client_secret") or not cfg.get("tenant_domain"):
        raise HTTPException(status_code=400, detail="Microsoft 365 non configuré.")
    removed, failed = [], []
    for email in data.emails:
        email = email.strip().lower()
        if not email:
            continue
        rule = f"SigChamp - {email}"
        try:
            # Action principale : retirer la règle de flux Exchange (signature serveur).
            r = await asyncio.to_thread(_exo_invoke, cfg, "Remove-TransportRule", {"Identity": rule, "Confirm": False})
            # Nettoyer aussi une éventuelle signature de compte posée par une ancienne version.
            try:
                await asyncio.to_thread(_exo_invoke, cfg, "Set-MailboxMessageConfiguration",
                                        {"Identity": email, "SignatureHtml": "", "SignatureText": " ",
                                         "AutoAddSignature": False, "AutoAddSignatureOnReply": False,
                                         "AutoAddSignatureOnMobile": False,
                                         "DefaultSignature": "", "DefaultSignatureOnReply": "",
                                         "DeleteSignatureName": "SigChamp"})
            except Exception:
                pass
            low = "".join(ch for ch in r.text if ch.isprintable()).lower()
            not_found = any(s in low for s in ["couldn't be found", "wasn't found", "n'existe", "not found", "objectnotfound"])
            if r.status_code < 300 or not_found:
                removed.append(email)
                emp = await db.users.find_one({"email": email})
                if emp:
                    await db.users.update_one({"_id": emp["_id"]}, {"$set": {"m365_pushed_at": None}})
            else:
                failed.append({"email": email, "error": _exo_error(r)})
        except HTTPException as e:
            failed.append({"email": email, "error": str(e.detail)})
        except Exception as e:
            failed.append({"email": email, "error": str(e)})
    return {"removed": removed, "failed": failed, "removed_count": len(removed)}


class AssetsConfig(BaseModel):
    enabled: bool = False
    provider: str = "ftp"  # "ftp" (SiteGround) ou "r2" (Cloudflare R2)
    backup_enabled: bool = False  # copier aussi vers SiteGround/FTP quand R2 est principal
    auto_cleanup: bool = False  # supprimer auto les orphelins R2 après changement de logo/photo
    public_url: Optional[str] = None
    ftp_host: Optional[str] = None
    ftp_port: Optional[int] = 21
    ftp_user: Optional[str] = None
    ftp_password: Optional[str] = None
    ftp_dir: Optional[str] = None
    ftp_secure: bool = True
    r2_account_id: Optional[str] = None
    r2_access_key: Optional[str] = None
    r2_secret_key: Optional[str] = None
    r2_bucket: Optional[str] = None


@api_router.get("/assets/config")
async def get_assets_config(admin: dict = Depends(require_perm("hosting"))):
    cfg = await db.config.find_one({"key": "assets"}) or {}
    return {"enabled": bool(cfg.get("enabled")), "provider": cfg.get("provider", "ftp"),
            "backup_enabled": bool(cfg.get("backup_enabled")),
            "auto_cleanup": bool(cfg.get("auto_cleanup")),
            "public_url": cfg.get("public_url", ""),
            "ftp_host": cfg.get("ftp_host", ""), "ftp_port": cfg.get("ftp_port", 21),
            "ftp_user": cfg.get("ftp_user", ""), "ftp_dir": cfg.get("ftp_dir", ""),
            "ftp_secure": bool(cfg.get("ftp_secure", True)), "has_password": bool(cfg.get("ftp_password")),
            "r2_account_id": cfg.get("r2_account_id", ""), "r2_access_key": cfg.get("r2_access_key", ""),
            "r2_bucket": cfg.get("r2_bucket", ""), "has_r2_secret": bool(cfg.get("r2_secret_key"))}


@api_router.put("/assets/config")
async def put_assets_config(data: AssetsConfig, admin: dict = Depends(require_perm("hosting"))):
    existing = await db.config.find_one({"key": "assets"}) or {}
    payload = {"key": "assets", "enabled": data.enabled, "provider": data.provider or "ftp",
               "backup_enabled": bool(data.backup_enabled),
               "auto_cleanup": bool(data.auto_cleanup),
               "public_url": (data.public_url or "").strip().rstrip("/"),
               "ftp_host": (data.ftp_host or "").strip(), "ftp_port": data.ftp_port or 21,
               "ftp_user": (data.ftp_user or "").strip(), "ftp_dir": (data.ftp_dir or "").strip(),
               "ftp_secure": data.ftp_secure,
               "r2_account_id": (data.r2_account_id or "").strip(),
               "r2_access_key": (data.r2_access_key or "").strip(),
               "r2_bucket": (data.r2_bucket or "").strip()}
    if data.ftp_password:
        payload["ftp_password"] = data.ftp_password
    elif existing.get("ftp_password"):
        payload["ftp_password"] = existing["ftp_password"]
    if data.r2_secret_key:
        payload["r2_secret_key"] = data.r2_secret_key
    elif existing.get("r2_secret_key"):
        payload["r2_secret_key"] = existing["r2_secret_key"]
    await db.config.update_one({"key": "assets"}, {"$set": payload}, upsert=True)
    await refresh_assets_cache()
    return {"ok": True}


@api_router.post("/assets/test")
async def test_assets(admin: dict = Depends(require_perm("hosting"))):
    cfg = await db.config.find_one({"key": "assets"}) or {}
    provider = cfg.get("provider") or "ftp"
    if provider == "r2":
        if not cfg.get("r2_account_id") or not cfg.get("r2_bucket") or not cfg.get("r2_access_key"):
            raise HTTPException(status_code=400, detail="Renseignez d'abord l'Account ID, la clé d'accès et le bucket R2.")
    elif not cfg.get("ftp_host") or not cfg.get("ftp_user"):
        raise HTTPException(status_code=400, detail="Renseignez d'abord l'hôte et l'utilisateur FTP.")

    def _run():
        res = {"primary": None, "backup": None}
        try:
            if provider == "r2":
                _r2_upload(cfg, f"{APP_NAME}/_sigchamp_test.txt", b"SigChamp connexion OK", "text/plain")
            else:
                _ftp_upload(cfg, f"{APP_NAME}/_sigchamp_test.txt", b"SigChamp connexion OK")
            res["primary"] = {"ok": True}
        except Exception as e:
            res["primary"] = {"ok": False, "error": str(e)}
        if provider == "r2" and cfg.get("backup_enabled") and cfg.get("ftp_host") and cfg.get("ftp_user"):
            try:
                _ftp_upload(cfg, f"{APP_NAME}/_sigchamp_test.txt", b"SigChamp backup OK")
                res["backup"] = {"ok": True}
            except Exception as e:
                res["backup"] = {"ok": False, "error": str(e)}
        return res

    res = await asyncio.to_thread(_run)
    if not res["primary"].get("ok"):
        label = "R2" if provider == "r2" else "FTP"
        raise HTTPException(status_code=400, detail=f"Connexion {label} échouée : {res['primary'].get('error')}")
    pub = (cfg.get("public_url") or "").rstrip("/")
    out = {"ok": True, "test_url": f"{pub}/{APP_NAME}/_sigchamp_test.txt" if pub else ""}
    if res["backup"] is not None:
        out["backup_ok"] = res["backup"].get("ok")
        if not res["backup"].get("ok"):
            out["backup_error"] = res["backup"].get("error")
    return out


@api_router.post("/assets/publish-all")
async def publish_all_assets(admin: dict = Depends(require_perm("hosting"))):
    cfg = await db.config.find_one({"key": "assets"}) or {}
    provider = cfg.get("provider") or "ftp"
    if provider == "r2":
        if not cfg.get("r2_bucket"):
            raise HTTPException(status_code=400, detail="Configurez d'abord Cloudflare R2.")
    elif not cfg.get("ftp_host"):
        raise HTTPException(status_code=400, detail="Configurez d'abord l'hébergement FTP.")

    def _run():
        count, failed = 0, []
        # 1) Tous les fichiers du stockage local (logos, GIF, avatars) en conservant les chemins.
        if STORAGE_BACKEND == "local" and LOCAL_STORAGE_DIR.exists():
            for fp in LOCAL_STORAGE_DIR.rglob("*"):
                if fp.is_file():
                    rel = str(fp.relative_to(LOCAL_STORAGE_DIR)).replace("\\", "/")
                    try:
                        _publish_asset(cfg, rel, fp.read_bytes())
                        count += 1
                    except Exception as e:
                        failed.append({"file": rel, "error": str(e)})
        # 2) Icônes des réseaux sociaux → dossier social/.
        social_dir = ASSETS_DIR / "social"
        if social_dir.exists():
            for fp in social_dir.glob("*.png"):
                try:
                    _publish_asset(cfg, f"social/{fp.name}", fp.read_bytes())
                    count += 1
                except Exception as e:
                    failed.append({"file": f"social/{fp.name}", "error": str(e)})
        return {"published": count, "failed": failed}

    return await asyncio.to_thread(_run)


async def _referenced_asset_keys() -> set:
    """Clés d'objets (chemins) actuellement utilisées par les signatures : logo, bannières GIF, avatars.
    Indépendant du domaine : on extrait la portion après '/api/files/' quel que soit l'hôte (évite les
    problèmes après un changement d'URL de déploiement)."""
    keys = set()
    marker = "/api/files/"

    def add(url):
        if not url:
            return
        u = str(url).split("?")[0]
        idx = u.find(marker)
        if idx != -1:
            keys.add(u[idx + len(marker):])

    s = await db.settings.find_one({"key": "global"}) or {}
    add(s.get("logo_url"))
    add(s.get("gif_url"))
    for b in s.get("department_banners", []):
        add(b.get("gif_url"))
    async for u in db.users.find({}, {"avatar_url": 1}):
        add(u.get("avatar_url"))
    return keys


@api_router.get("/assets/usage")
async def assets_usage(admin: dict = Depends(require_perm("hosting"))):
    """Usage du stockage R2 (Go utilisés / quota gratuit 10 Go) + nombre de fichiers orphelins."""
    cfg = await db.config.find_one({"key": "assets"}) or {}
    if (cfg.get("provider") or "ftp") != "r2" or not cfg.get("r2_bucket") or not cfg.get("r2_access_key"):
        return {"available": False, "reason": "R2 n'est pas configuré comme fournisseur principal."}
    referenced = await _referenced_asset_keys()

    def _run():
        try:
            objs = _r2_list(cfg)
            total = sum(o["size"] for o in objs)
            orphans = [o for o in objs if not o["key"].startswith("social/") and o["key"] not in referenced]
            return {"ok": True, "count": len(objs), "bytes": total,
                    "orphan_count": len(orphans), "orphan_bytes": sum(o["size"] for o in orphans)}
        except Exception as e:
            return {"ok": False, "error": str(e)}

    r = await asyncio.to_thread(_run)
    if not r.get("ok"):
        raise HTTPException(status_code=400, detail=f"Lecture de l'usage R2 échouée : {r.get('error')}")
    quota = 10 * 1024 ** 3
    return {"available": True, "used_bytes": r["bytes"], "quota_bytes": quota,
            "used_gb": round(r["bytes"] / 1024 ** 3, 4), "quota_gb": 10,
            "percent": round(r["bytes"] / quota * 100, 2), "object_count": r["count"],
            "orphan_count": r["orphan_count"], "orphan_bytes": r["orphan_bytes"]}


async def _cleanup_orphans_core() -> dict:
    """Supprime les fichiers orphelins (non référencés) de R2, du backup FTP et du stockage local."""
    cfg = await db.config.find_one({"key": "assets"}) or {}
    provider = cfg.get("provider") or "ftp"
    referenced = await _referenced_asset_keys()

    # Garde-fou : si aucune référence n'est détectée, on n'efface RIEN (évite une suppression massive
    # en cas d'échec de détection, ex. base vide ou incohérente).
    if not referenced:
        return {"deleted": 0, "freed_bytes": 0, "failed": [],
                "skipped": "Aucune image référencée détectée — nettoyage annulé par sécurité."}

    def _run():
        deleted, freed, failed = 0, 0, []
        orphan_keys = []
        if provider == "r2" and cfg.get("r2_bucket"):
            try:
                objs = _r2_list(cfg)
                orphan_keys = [o["key"] for o in objs
                               if not o["key"].startswith("social/") and o["key"] not in referenced]
                freed = sum(o["size"] for o in objs if o["key"] in set(orphan_keys))
                _r2_delete(cfg, orphan_keys)
                deleted = len(orphan_keys)
            except Exception as e:
                failed.append({"target": "r2", "error": str(e)})
        # Backup FTP : supprimer les mêmes clés orphelines (best-effort).
        if cfg.get("backup_enabled") and cfg.get("ftp_host") and cfg.get("ftp_user"):
            for k in orphan_keys:
                try:
                    _ftp_delete(cfg, k)
                except Exception:
                    pass
        # Stockage local : supprimer les fichiers non référencés (garde uniquement les fichiers actuels).
        if STORAGE_BACKEND == "local" and LOCAL_STORAGE_DIR.exists():
            for fp in LOCAL_STORAGE_DIR.rglob("*"):
                if fp.is_file():
                    rel = str(fp.relative_to(LOCAL_STORAGE_DIR)).replace("\\", "/")
                    if rel not in referenced:
                        try:
                            fp.unlink()
                        except Exception:
                            pass
        return {"deleted": deleted, "freed_bytes": freed, "failed": failed}

    return await asyncio.to_thread(_run)


async def _maybe_auto_cleanup():
    """Lance le nettoyage des orphelins en arrière-plan si l'auto-nettoyage R2 est activé."""
    cfg = _assets_cache or {}
    if cfg.get("enabled") and (cfg.get("provider") or "ftp") == "r2" and cfg.get("auto_cleanup"):
        async def _bg():
            try:
                res = await _cleanup_orphans_core()
                if res.get("deleted"):
                    logger.info(f"Auto-nettoyage R2 : {res['deleted']} orphelin(s) supprimé(s), {res['freed_bytes']} octet(s) libéré(s).")
            except Exception as e:
                logger.warning(f"Auto-nettoyage R2 échoué : {e}")
        asyncio.create_task(_bg())


@api_router.post("/assets/cleanup")
async def assets_cleanup(admin: dict = Depends(require_perm("hosting"))):
    """Supprime les fichiers orphelins (non référencés par les signatures actuelles) de R2, du backup FTP et du stockage local."""
    return await _cleanup_orphans_core()



@api_router.get("/system/status")
async def system_status(admin: dict = Depends(require_perm("system"))):
    services = []
    _hidden = {"code-server", "nginx-code-proxy", "webhook-crond"}
    try:
        out = subprocess.run(["sudo", "supervisorctl", "status"], capture_output=True, text=True, timeout=15).stdout
        for line in out.strip().splitlines():
            parts = line.split()
            if len(parts) >= 2 and parts[0] not in _hidden:
                services.append({"name": parts[0], "state": parts[1],
                                 "info": " ".join(parts[2:])})
    except Exception as e:
        return {"services": [], "error": str(e), "restart_command": _RESTART_COMMAND}
    return {"services": services, "restart_command": _RESTART_COMMAND}


@api_router.post("/system/reload")
async def system_reload(admin: dict = Depends(require_perm("system"))):
    """Rechargement léger : relit la configuration (.env) et vérifie les données, sans couper le service."""
    load_dotenv(ROOT_DIR / ".env", override=True)
    await load_settings()
    await db.command("ping")
    return {"ok": True, "message": "Configuration rechargée et connexion à la base vérifiée. Aucune interruption de service."}


@api_router.post("/system/restart")
async def system_restart(admin: dict = Depends(require_perm("system"))):
    """Redémarrage complet des services backend + frontend. Fonctionne dans l'aperçu (supervisor)
    et en auto-hébergé (définir RESTART_COMMAND, ex. 'docker compose restart')."""
    def _do_restart():
        try:
            subprocess.Popen(f"sleep 1 && {_RESTART_COMMAND}", shell=True, start_new_session=True)
        except Exception as e:
            logger.error(f"Restart failed: {e}")
    _do_restart()
    return {"ok": True, "message": "Redémarrage des services lancé. Le service sera de nouveau disponible dans quelques secondes."}


@api_router.get("/")
async def root():
    return {"message": "SigChamp API"}


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def startup():
    if STORAGE_BACKEND == "local":
        LOCAL_STORAGE_DIR.mkdir(parents=True, exist_ok=True)
        logger.info(f"Storage: local disk at {LOCAL_STORAGE_DIR}")
    else:
        try:
            init_storage()
            logger.info("Storage: Emergent object storage")
        except Exception as e:
            logger.error(f"Storage init failed: {e}")

    await db.users.create_index("email", unique=True)
    admin_email = os.environ.get("ADMIN_EMAIL", "admin@sigflow.com").lower()
    admin_password = os.environ.get("ADMIN_PASSWORD", "admin123")
    existing = await db.users.find_one({"email": admin_email})
    if existing is None:
        await db.users.insert_one({
            "email": admin_email, "password_hash": hash_password(admin_password),
            "name": "Administrateur", "role": "admin", "title": "Gestionnaire de marque",
            "phone_ext": "", "direct_line": "", "department": "Direction", "avatar_url": "",
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        logger.info("Admin seeded")
    elif not verify_password(admin_password, existing["password_hash"]):
        await db.users.update_one({"email": admin_email},
                                  {"$set": {"password_hash": hash_password(admin_password)}})
    await load_settings()
    await refresh_assets_cache()


@app.on_event("shutdown")
async def shutdown():
    client.close()
